# ==============================================================================
# kwaliteit_engine.py — Data Quality analyse (7 Dimensies)
# ==============================================================================

# --- Stdlib imports ---
import os
import time
from datetime import datetime, timezone

# --- Third-party imports ---
import docx
import openpyxl
import pptx
import PyPDF2


# ==============================================================================
# QUALITY ENGINE
# ==============================================================================

class QualityEngine:
    """
    Analyseert de technische kwaliteit van bestanden aan de hand van de
    7 DAMA DMBOK dimensies (vertaald naar documenten).
    """

    def __init__(self, active_modules: list[str]):
        self.active_rules = active_modules
        
        # Mapping van rules naar overkoepelende dimensies voor de score aggregatie
        self.dimension_mapping = {
            "Auteurs-validatie": "1. Nauwkeurigheid (Accuracy)",
            "Bestandsbody Check": "2. Volledigheid (Completeness)",
            "Metadata (Titel) Check": "2. Volledigheid (Completeness)",
            "Extensie-correlatie": "3. Consistentie (Consistency)",
            "Actualiteits-norm": "4. Tijdigheid (Timeliness)",
            "Leesbaarheids-garantie": "5. Validiteit (Validity)",
            "Geen vreemde tekens": "5. Validiteit (Validity)",
            "Exacte duplicatie": "6. Uniciteit (Uniqueness)",
            "Dode Snelkoppelingen": "7. Integriteit (Integrity)"
        }
        
        # De uiteindelijke domeinen die het dashboard verwacht:
        self.domains = list(set([self.dimension_mapping[rule] for rule in self.active_rules if rule in self.dimension_mapping]))

    def analyze(self, item: dict, stream) -> dict:
        scores_per_dim = {dim: [] for dim in self.domains}
        reden = []

        filename = item["name"].lower()
        extension = item["extension"]
        is_readable_doc = item.get("is_readable_doc", False)
        file_is_locked = stream is None

        # --- 1. Nauwkeurigheid ---
        if "Auteurs-validatie" in self.active_rules:
            dim = self.dimension_mapping["Auteurs-validatie"]
            if file_is_locked or not is_readable_doc:
                scores_per_dim[dim].append("N/A")
            else:
                has_author = self._has_property(stream, extension, 'author')
                if has_author:
                    scores_per_dim[dim].append(100)
                else:
                    scores_per_dim[dim].append(0)
                    reden.append("Nauwkeurigheid: Auteur onbekend in document-eigenschappen.")

        # --- 2. Volledigheid ---
        if "Bestandsbody Check" in self.active_rules:
            dim = self.dimension_mapping["Bestandsbody Check"]
            if item.get("size", 0) < 1024:
                scores_per_dim[dim].append(0)
                reden.append("Volledigheid: Bestand is vrijwel leeg (<1KB).")
            else:
                scores_per_dim[dim].append(100)
                
        if "Metadata (Titel) Check" in self.active_rules:
            dim = self.dimension_mapping["Metadata (Titel) Check"]
            if file_is_locked or not is_readable_doc:
                scores_per_dim[dim].append("N/A")
            else:
                has_title = self._has_property(stream, extension, 'title')
                if has_title:
                    scores_per_dim[dim].append(100)
                else:
                    scores_per_dim[dim].append(0)
                    reden.append("Volledigheid: Document heeft geen ingevulde 'Titel' eigenschap.")

        # --- 3. Consistentie ---
        if "Extensie-correlatie" in self.active_rules:
            dim = self.dimension_mapping["Extensie-correlatie"]
            if extension in ['.exe', '.bat', '.sh', '', '.dll']:
                scores_per_dim[dim].append(0)
                reden.append(f"Consistentie: Ongeldig document-formaat ({extension}).")
            else:
                scores_per_dim[dim].append(100)

        # --- 4. Tijdigheid ---
        if "Actualiteits-norm" in self.active_rules:
            dim = self.dimension_mapping["Actualiteits-norm"]
            age = self._calculate_age(item)
            if age > 5:
                scores_per_dim[dim].append(0)
                reden.append(f"Tijdigheid: Zeer oud archiefbestand ({age:.1f} jr).")
            elif age > 3:
                scores_per_dim[dim].append(50)
                reden.append(f"Tijdigheid: Ouder dan 3 jaar ({age:.1f} jr).")
            else:
                scores_per_dim[dim].append(100)

        # --- 5. Validiteit ---
        if "Leesbaarheids-garantie" in self.active_rules:
            dim = self.dimension_mapping["Leesbaarheids-garantie"]
            if file_is_locked and item.get("size", 0) > 0:
                scores_per_dim[dim].append(0)
                reden.append("Validiteit: Bestand is vergrendeld of corrupt.")
            else:
                scores_per_dim[dim].append(100)
                
        if "Geen vreemde tekens" in self.active_rules:
            dim = self.dimension_mapping["Geen vreemde tekens"]
            if item.get("has_forbidden_chars", False):
                scores_per_dim[dim].append(0)
                reden.append("Validiteit: Bestandsnaam bevat illegale tekens.")
            else:
                scores_per_dim[dim].append(100)

        # --- 6. Uniciteit ---
        if "Exacte duplicatie" in self.active_rules:
            dim = self.dimension_mapping["Exacte duplicatie"]
            if item.get("is_duplicate", False):
                scores_per_dim[dim].append(0)
                reden.append("Uniciteit: Systeem heeft dubbele kopieën gevonden.")
            else:
                scores_per_dim[dim].append(100)

        # --- 7. Integriteit ---
        if "Dode Snelkoppelingen" in self.active_rules:
            dim = self.dimension_mapping["Dode Snelkoppelingen"]
            if extension in ['.lnk', '.url']:
                scores_per_dim[dim].append(0)
                reden.append("Integriteit: Snelkoppeling naar mogelijk dode/onveilige link.")
            else:
                scores_per_dim[dim].append(100)

        # Resultaten integreren
        final_scores = {}
        for dim, s_list in scores_per_dim.items():
            valid_scores = [s for s in s_list if isinstance(s, int)]
            if not valid_scores:
                final_scores[dim] = "N/A"
            else:
                final_scores[dim] = int(sum(valid_scores) / len(valid_scores))

        return {"scores": final_scores, "reasons": list(set(reden))}

    # --- Helpers ---
    def _calculate_age(self, item: dict) -> float:
        try:
            now = time.time()
            if item["mode"] == "local":
                mtime = item.get("time_modified", os.path.getmtime(item["path"]))
                return (now - mtime) / (365 * 24 * 3600)
            elif item["mode"] == "sp":
                ds = str(item.get("time_modified", ""))
                if "T" in ds and "Z" in ds:
                    dt = datetime.strptime(ds, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
                    return (now - dt.timestamp()) / (365 * 24 * 3600)
            return 0
        except Exception:
            return 0

    def _has_property(self, stream, ext: str, prop: str) -> bool:
        if not stream or ext not in ['.pdf', '.docx', '.xlsx', '.pptx']:
            return False
        try:
            stream.seek(0)
            if ext == '.pdf':
                meta = PyPDF2.PdfReader(stream).metadata
                return bool(meta and prop in str(meta).lower())
            elif ext == '.docx':
                if prop == 'author': return bool(docx.Document(stream).core_properties.author)
                if prop == 'title': return bool(docx.Document(stream).core_properties.title)
            elif ext == '.xlsx':
                wb = openpyxl.load_workbook(stream, read_only=True)
                if prop == 'author': return bool(wb.properties.creator)
                if prop == 'title': return bool(wb.properties.title)
            elif ext == '.pptx':
                prs = pptx.Presentation(stream)
                if prop == 'author': return bool(prs.core_properties.author)
                if prop == 'title': return bool(prs.core_properties.title)
        except Exception:
            pass
        return False
